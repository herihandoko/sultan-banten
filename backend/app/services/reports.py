"""Report builders — Excel (openpyxl) + PDF (fpdf2) for PRD §9.1."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from fpdf import FPDF
from fpdf.fonts import FontFace

from app.extensions import db
from app.models import (
    ContentItem,
    Issue,
    KolCampaign,
    KolPartner,
    MediaBlastLog,
    MediaPartner,
    MediaSlaLog,
    Mission,
    MissionParticipation,
    OpdValidation,
    Report,
)
from app.models.issue import public_actions
from sqlalchemy import func


def _now():
    return datetime.now(timezone.utc)


def _wb_bytes(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]):
    ws = wb.active if wb.active.title == "Sheet" and not wb.active["A1"].value else wb.create_sheet(title)
    if wb.active.title == "Sheet" and not wb.active["A1"].value:
        ws.title = title
    ws.append(headers)
    for row in rows:
        ws.append(row)
    return ws


class _ReportPDF(FPDF):
    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", size=8)
        self.set_text_color(120, 120, 120)
        self.cell(0, 8, f"SIAGAPIM - halaman {self.page_no()}", align="C")


def _safe(text: str) -> str:
    """Helvetica core fonts are Latin-1; normalize common punctuation first."""
    if text is None:
        return ""
    s = str(text)
    for src, dst in (
        ("\u2014", "-"),
        ("\u2013", "-"),
        ("\u2018", "'"),
        ("\u2019", "'"),
        ("\u201c", '"'),
        ("\u201d", '"'),
        ("\u2026", "..."),
        ("\u00a0", " "),
    ):
        s = s.replace(src, dst)
    return s.encode("latin-1", errors="replace").decode("latin-1")


def _pdf_bytes(title: str, lines: list[str], sections: list[tuple[str, list[str]]] | None = None) -> bytes:
    pdf = _ReportPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 8, _safe(title))
    pdf.set_font("Helvetica", size=10)
    pdf.set_text_color(80, 80, 80)
    pdf.cell(0, 6, f"Dibuat: {_now().strftime('%Y-%m-%d %H:%M UTC')}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    pdf.set_text_color(20, 20, 20)
    for line in lines:
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", size=11)
        pdf.multi_cell(0, 6, _safe(line))
    for section_title, section_lines in sections or []:
        pdf.ln(3)
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "B", 12)
        pdf.multi_cell(0, 7, _safe(section_title))
        pdf.set_font("Helvetica", size=10)
        for line in section_lines:
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(0, 5, _safe(f"- {line}"))
    out = pdf.output()
    if isinstance(out, (bytes, bytearray)):
        return bytes(out)
    return out.encode("latin-1")


def _save_report(report_type: str, title: str, payload: dict, user_id: int | None) -> Report:
    report = Report(
        report_type=report_type,
        title=title,
        period_end=_now(),
        payload=payload,
        generated_by=user_id,
    )
    db.session.add(report)
    db.session.commit()
    return report


# ── Crisis ──────────────────────────────────────────────────────────────────


def collect_crisis(issue_id: int | None = None, project_id: str | None = None) -> dict:
    from app.utils.project_scope import filter_issues_query

    q = Issue.query.order_by(Issue.created_at.desc())
    q = filter_issues_query(q, project_id)
    if issue_id:
        q = q.filter_by(id=issue_id)
    issues = q.limit(50).all()
    rows = []
    for issue in issues:
        vals = OpdValidation.query.filter_by(issue_id=issue.id).all()
        contents = ContentItem.query.filter_by(issue_id=issue.id).all()
        blasts = MediaBlastLog.query.filter_by(issue_id=issue.id).count()
        status_label, actor = _STATUS_LABEL.get(issue.status, (issue.status or "-", "-"))
        rows.append(
            {
                "id": issue.id,
                "title": issue.title,
                "risk_level": issue.risk_level,
                "risk_label": _lbl(_RISK_LABEL, issue.risk_level),
                "status": issue.status,
                "status_label": status_label,
                "actor": actor,
                "source": issue.primary_source_label() or "Tidak tercatat",
                "mb_alert_id": issue.mb_alert_id,
                "summary": (issue.summary or "").strip(),
                "validations": len(vals),
                "validated": sum(1 for v in vals if v.status == "validated"),
                "contents": len(contents),
                "approved": sum(1 for c in contents if c.status in {"approved", "published"}),
                "blasts": blasts,
                "created_at": issue.created_at.isoformat() if issue.created_at else None,
                "created_label": _fmt_wib(issue.created_at),
                "closed_at": issue.closed_at.isoformat() if issue.closed_at else None,
            }
        )
    return {"type": "crisis", "count": len(rows), "issues": rows}


_NAVY = (27, 58, 92)
_GOLD = (176, 138, 52)
_INK = (28, 32, 38)
_MUTED = (90, 98, 110)

_MONTHS_ID = (
    "Januari",
    "Februari",
    "Maret",
    "April",
    "Mei",
    "Juni",
    "Juli",
    "Agustus",
    "September",
    "Oktober",
    "November",
    "Desember",
)

_RISK_LABEL = {
    "R0": "R0 - Informasional",
    "R1": "R1 - Observasi",
    "R2": "R2 - Waspada",
    "R3": "R3 - Meningkat",
    "R4": "R4 - Tinggi",
    "R5": "R5 - Kritis",
}

_STATUS_LABEL = {
    "open": ("Terbuka", "Editor"),
    "validating": ("Dalam validasi", "OPD"),
    "producing": ("Produksi konten", "Editor"),
    "approved": ("Disetujui", "Pimpinan"),
    "disseminated": ("Telah disebarluaskan", "Media"),
    "closed": ("Ditutup", "Editor"),
}

_VALIDATION_LABEL = {
    "waiting": "Menunggu tanggapan",
    "validated": "Tervalidasi",
    "rejected": "Ditolak",
}

_CONTENT_STATUS = {
    "draft": "Draf",
    "in_review": "Menunggu persetujuan",
    "approved": "Disetujui",
    "rejected": "Ditolak",
    "published": "Telah diterbitkan",
}

_CONTENT_TYPE = {
    "text_release": "Rilis teks",
    "infographic": "Infografis",
    "video": "Video",
}

_MISSION_STATUS = {
    "active": "Berjalan",
    "completed": "Selesai",
    "cancelled": "Dibatalkan",
}

_ACTION_LABEL = {
    "like": "Menyukai unggahan",
    "share": "Membagikan unggahan",
    "comment": "Memberi komentar",
    "like_share_comment": "Menyukai, membagikan, dan mengomentari",
}

_CAMPAIGN_STATUS = {
    "planned": "Direncanakan",
    "in_progress": "Berjalan",
    "published": "Telah tayang",
    "completed": "Selesai",
    "cancelled": "Dibatalkan",
}

_CONTRACT_LABEL = {
    "prospect": "Prospek",
    "active": "Aktif",
    "expired": "Habis masa berlaku",
    "terminated": "Dihentikan",
}

_BUDGET_LABEL = {
    "planned": "Anggaran rencana",
    "approved": "Anggaran disetujui",
    "paid": "Sudah dibayar",
    "cancelled": "Anggaran dibatalkan",
}


def _lbl(mapping: dict, key, fallback: str = "-") -> str:
    if key in mapping:
        return mapping[key]
    return str(key) if key else fallback


def _fmt_int(value) -> str:
    return f"{int(value or 0):,}".replace(",", ".")


def _fmt_idr(value) -> str:
    if value is None:
        return "-"
    return "Rp " + _fmt_int(round(float(value)))


def _fmt_pct(part, whole) -> str:
    if not whole:
        return "0%"
    text = f"{round(100 * part / whole, 1):.1f}"
    return text.replace(".", ",") + "%"


def _printer_name(user_id: int | None) -> str:
    if not user_id:
        return "Sistem SIAGAPIM"
    from app.models import User

    user = db.session.get(User, user_id)
    if not user:
        return "Sistem SIAGAPIM"
    role = user.role.name if user.role else ""
    return f"{user.full_name}" + (f", {role}" if role else "")


def _fmt_wib(dt: datetime | None) -> str:
    if not dt:
        return "-"
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    from zoneinfo import ZoneInfo

    local = dt.astimezone(ZoneInfo("Asia/Jakarta"))
    return f"{local.day} {_MONTHS_ID[local.month - 1]} {local.year}, {local.hour:02d}.{local.minute:02d} WIB"


def _clean_narrative(text: str | None) -> str:
    raw = (text or "").strip()
    if not raw:
        return ""
    lowered = raw.lower()
    if "deteksi otomatis" in lowered or "sipantau" in lowered or "mata bathin" in lowered:
        return ""
    return raw


class _IssueBriefPDF(FPDF):
    def header(self):
        if self.page_no() == 1:
            self.set_fill_color(*_NAVY)
            self.rect(0, 0, 210, 26, "F")
            self.set_fill_color(*_GOLD)
            self.rect(0, 26, 210, 1.4, "F")
            self.set_xy(16, 6)
            self.set_text_color(255, 255, 255)
            self.set_font("Helvetica", "B", 11)
            self.cell(0, 5, "PEMERINTAH PROVINSI BANTEN", new_x="LMARGIN", new_y="NEXT")
            self.set_x(16)
            self.set_font("Helvetica", size=8)
            self.cell(
                0,
                4,
                "SIAGAPIM  |  Sistem Informasi Analisis Gema Aktual Pimpinan",
                new_x="LMARGIN",
                new_y="NEXT",
            )
            self.set_x(16)
            self.set_font("Helvetica", size=7)
            self.set_text_color(210, 220, 230)
            self.cell(0, 4, "Biro Administrasi Pimpinan  -  Dokumen internal", new_x="LMARGIN", new_y="NEXT")
            self.set_y(34)
        else:
            self.set_font("Helvetica", size=8)
            self.set_text_color(*_MUTED)
            running = getattr(self, "running_title", "Laporan Penanganan Isu")
            self.cell(0, 6, _safe(f"SIAGAPIM Banten  -  {running}"), new_x="LMARGIN", new_y="NEXT")
            self.set_draw_color(*_NAVY)
            self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
            self.ln(3)

    def footer(self):
        self.set_y(-12)
        self.set_draw_color(210, 214, 220)
        self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
        self.ln(1.5)
        self.set_font("Helvetica", size=7)
        self.set_text_color(*_MUTED)
        width = self.w - self.l_margin - self.r_margin
        self.set_x(self.l_margin)
        self.cell(width * 0.72, 4, "Dokumen internal. Tidak untuk disebarluaskan di luar keperluan dinas.")
        self.cell(width * 0.28, 4, f"Halaman {self.page_no()}/{{nb}}", align="R")


def _brief_heading(pdf: FPDF, number: str, title: str) -> None:
    pdf.ln(3)
    pdf.set_x(pdf.l_margin)
    pdf.set_fill_color(*_NAVY)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 10)
    pdf.cell(0, 7, _safe(f"   {number}.   {title}"), fill=True, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_text_color(*_INK)


def _brief_label_value(pdf: FPDF, label: str, value: str) -> None:
    pdf.set_x(pdf.l_margin)
    pdf.set_font("Helvetica", "B", 9)
    pdf.set_text_color(*_MUTED)
    pdf.cell(46, 5.5, _safe(label))
    pdf.set_font("Helvetica", size=10)
    pdf.set_text_color(*_INK)
    pdf.multi_cell(0, 5.5, _safe(value or "-"))


def _brief_paragraph(pdf: FPDF, text: str) -> None:
    pdf.set_x(pdf.l_margin)
    pdf.set_font("Helvetica", size=10)
    pdf.set_text_color(*_INK)
    pdf.multi_cell(0, 5.4, _safe(text or "Tidak tersedia."))
    pdf.ln(1)


def _pdf_table(pdf: FPDF, headers: list[str], rows: list[list], widths: tuple[int, ...]) -> None:
    if not rows:
        _brief_paragraph(pdf, "Tidak ada data pada bagian ini.")
        return
    pdf.set_font("Helvetica", size=8)
    pdf.set_text_color(*_INK)
    pdf.set_fill_color(255, 255, 255)
    usable = pdf.w - pdf.l_margin - pdf.r_margin
    with pdf.table(
        width=usable,
        col_widths=widths,
        line_height=4.6,
        text_align="LEFT",
        headings_style=FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=_NAVY),
    ) as table:
        head = table.row()
        for cell in headers:
            head.cell(_safe(cell))
        for data in rows:
            row = table.row()
            for cell in data:
                row.cell(_safe("-" if cell is None or cell == "" else str(cell)))
    pdf.ln(2)


def _official_pdf(running_title: str, title: str, doc_no: str, printed_by: str, blocks: list) -> bytes:
    pdf = _IssueBriefPDF()
    pdf.running_title = running_title
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(16, 16, 16)
    pdf.add_page()
    printed = _fmt_wib(_now())
    pdf.set_text_color(*_NAVY)
    pdf.set_font("Helvetica", "B", 14)
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(0, 7, _safe(title))
    pdf.set_font("Helvetica", size=9)
    pdf.set_text_color(*_MUTED)
    pdf.set_x(pdf.l_margin)
    pdf.cell(0, 5, _safe(f"Nomor: {doc_no}"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_x(pdf.l_margin)
    pdf.cell(0, 5, _safe(f"Dicetak: {printed}"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_x(pdf.l_margin)
    pdf.cell(0, 5, _safe(f"Penyusun: {printed_by}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(1)
    for block in blocks:
        kind = block[0]
        if kind == "h":
            _brief_heading(pdf, block[1], block[2])
        elif kind == "kv":
            for label, value in block[1]:
                _brief_label_value(pdf, label, str(value))
        elif kind == "p":
            _brief_paragraph(pdf, block[1])
        elif kind == "items":
            if not block[1]:
                _brief_paragraph(pdf, "Tidak ada data pada bagian ini.")
            for line in block[1]:
                pdf.set_x(pdf.l_margin)
                pdf.set_font("Helvetica", size=10)
                pdf.set_text_color(*_INK)
                pdf.multi_cell(0, 5.2, _safe(line))
                pdf.ln(0.6)
        elif kind == "table":
            _pdf_table(pdf, block[1], block[2], block[3])
        elif kind == "close":
            pdf.ln(3)
            pdf.set_x(pdf.l_margin)
            pdf.set_font("Helvetica", "I", 9)
            pdf.set_text_color(*_MUTED)
            pdf.multi_cell(0, 4.8, _safe(block[1]))
    out = pdf.output()
    return bytes(out) if isinstance(out, (bytes, bytearray)) else out.encode("latin-1")


_THIN = Border(
    left=Side(style="thin", color="D0D7E2"),
    right=Side(style="thin", color="D0D7E2"),
    top=Side(style="thin", color="D0D7E2"),
    bottom=Side(style="thin", color="D0D7E2"),
)


def _write_cover(ws, cover: dict) -> None:
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.7, right=0.6, top=0.6, bottom=0.6)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 78
    ws.merge_cells("A1:B1")
    ws["A1"] = "PEMERINTAH PROVINSI BANTEN"
    ws["A1"].font = Font(name="Calibri", bold=True, size=16, color="1B3A5C")
    ws.merge_cells("A2:B2")
    ws["A2"] = "Biro Administrasi Pimpinan"
    ws["A2"].font = Font(name="Calibri", size=12, color="1B3A5C")
    ws.merge_cells("A3:B3")
    ws["A3"] = "SIAGAPIM — Sistem Informasi Analisis Gema Aktual Pimpinan"
    ws["A3"].font = Font(name="Calibri", size=10, italic=True, color="5A626E")
    ws.merge_cells("A5:B5")
    ws["A5"] = cover["title"]
    ws["A5"].font = Font(name="Calibri", bold=True, size=14, color="1B3A5C")
    ws["A5"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[5].height = 24
    meta = [
        (6, "Nomor", cover["doc_no"]),
        (7, "Tanggal cetak", cover["printed_at"]),
        (8, "Penyusun", cover["printed_by"]),
        (9, "Perihal", cover["subject"]),
    ]
    for row, label, value in meta:
        ws.cell(row, 1, label).font = Font(name="Calibri", bold=True, size=11, color="5A626E")
        cell = ws.cell(row, 2, value)
        cell.font = Font(name="Calibri", size=11, color="1C2026")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws["A11"] = "IKHTISAR"
    ws["A11"].font = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
    ws["A11"].fill = PatternFill("solid", fgColor="1B3A5C")
    ws["B11"].fill = PatternFill("solid", fgColor="1B3A5C")
    alt = PatternFill("solid", fgColor="F4F7FB")
    for offset, (label, value) in enumerate(cover["facts"]):
        row = 12 + offset
        left = ws.cell(row, 1, label)
        right = ws.cell(row, 2, value)
        for cell in (left, right):
            cell.font = Font(name="Calibri", size=11, color="1C2026")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = _THIN
            if offset % 2 == 0:
                cell.fill = alt
        left.font = Font(name="Calibri", bold=True, size=11, color="1B3A5C")
        ws.row_dimensions[row].height = 20
    note_row = 13 + len(cover["facts"])
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=2)
    ws.cell(note_row, 1, "URAIAN")
    ws.cell(note_row, 1).font = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
    ws.cell(note_row, 1).fill = PatternFill("solid", fgColor="1B3A5C")
    ws.cell(note_row, 2).fill = PatternFill("solid", fgColor="1B3A5C")
    body_row = note_row + 1
    ws.merge_cells(start_row=body_row, start_column=1, end_row=body_row + 4, end_column=2)
    body = ws.cell(body_row, 1, cover["narrative"])
    body.font = Font(name="Calibri", size=11, color="1C2026")
    body.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[body_row].height = 36
    foot = body_row + 6
    ws.merge_cells(start_row=foot, start_column=1, end_row=foot, end_column=2)
    ws.cell(foot, 1, "Dokumen internal. Tidak untuk disebarluaskan di luar keperluan dinas.")
    ws.cell(foot, 1).font = Font(name="Calibri", italic=True, size=9, color="5A626E")
    ws.oddHeader.left.text = "Pemerintah Provinsi Banten"
    ws.oddFooter.left.text = "Dokumen internal SIAGAPIM"
    ws.oddFooter.right.text = "Halaman &P"
    ws.print_title_rows = "1:3"
    ws.sheet_properties.tabColor = "1B3A5C"


def _write_table(ws, headers: list[str], rows: list[list], widths: list[int] | None = None) -> None:
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.horizontalCentered = True
    ws.page_margins = PageMargins(left=0.45, right=0.45, top=0.7, bottom=0.55, header=0.28, footer=0.28)
    ws.oddHeader.left.text = "SIAGAPIM — Pemerintah Provinsi Banten"
    ws.oddFooter.left.text = "Dokumen internal"
    ws.oddFooter.right.text = "Halaman &P dari &N"
    header_fill = PatternFill("solid", fgColor="1B3A5C")
    alt = PatternFill("solid", fgColor="F4F7FB")
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(1, col, header)
        cell.fill = header_fill
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 24
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions if rows else f"A1:{get_column_letter(len(headers))}1"
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(r_idx, c_idx, value if value is not None else "")
            cell.font = Font(name="Calibri", size=10, color="1C2026")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = _THIN
            if r_idx % 2 == 0:
                cell.fill = alt
        ws.row_dimensions[r_idx].height = 32
    for col, header in enumerate(headers, start=1):
        letter = get_column_letter(col)
        if widths and col - 1 < len(widths):
            ws.column_dimensions[letter].width = widths[col - 1]
            continue
        longest = len(str(header))
        for row in rows[:30]:
            if col - 1 < len(row):
                longest = max(longest, min(len(str(row[col - 1] or "")), 48))
        ws.column_dimensions[letter].width = min(max(longest + 3, 14), 46)
    ws.print_title_rows = "1:1"
    ws.page_setup.scale = 80
    if not rows:
        ws.cell(2, 1, "Tidak ada data pada periode ini.")
        ws.cell(2, 1).font = Font(name="Calibri", italic=True, size=10, color="5A626E")


def _formal_workbook(cover: dict, tables: list[dict]) -> bytes:
    wb = Workbook()
    wb.properties.title = cover["title"]
    wb.properties.creator = "SIAGAPIM Banten"
    wb.properties.subject = cover["subject"]
    wb.properties.category = "Laporan dinas"
    _write_cover(wb.active, cover)
    wb.active.title = "Ringkasan"
    for table in tables:
        ws = wb.create_sheet(table["name"][:31])
        _write_table(ws, table["headers"], table["rows"], table.get("widths"))
    return _wb_bytes(wb)


def _export_issue_brief(issue_id: int, user_id: int | None, project_id: str | None):
    from app.utils.project_scope import filter_issues_query

    query = filter_issues_query(Issue.query.filter_by(id=issue_id), project_id)
    issue = query.first()
    title = "Laporan Penanganan Isu"
    if not issue:
        pdf_bytes = _pdf_bytes(title, ["Isu tidak ditemukan atau berada di luar lingkup proyek yang dipilih."])
        return pdf_bytes, "application/pdf", f"laporan-isu-{issue_id}.pdf"

    status_label, actor = _STATUS_LABEL.get(issue.status, (issue.status or "-", "-"))
    risk_label = _RISK_LABEL.get(issue.risk_level, issue.risk_level or "-")
    printed = _fmt_wib(_now())
    year = _now().year
    doc_no = f"SIAGAPIM/ISU/{issue.id:04d}/{year}"

    vals = list(issue.validations or [])
    contents = list(issue.content_items or [])
    evidence = list(issue.evidence or [])
    blasts = MediaBlastLog.query.filter_by(issue_id=issue.id).count()
    outlet = issue.primary_source_label() or "-"
    source_url = issue.primary_source_url() or ""
    summary = (issue.summary or "").strip()
    why = _clean_narrative(issue.why_now)
    actions = public_actions(issue.recommended_actions) or []
    narrative = issue.narrative_card if isinstance(issue.narrative_card, dict) else {}
    statement = (narrative.get("statement") or "").strip()

    payload = {
        "type": "crisis_issue",
        "id": issue.id,
        "title": issue.title,
        "status": issue.status,
        "risk_level": issue.risk_level,
    }
    _save_report("crisis", title, payload, user_id)

    pdf = _IssueBriefPDF()
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(16, 16, 16)
    pdf.add_page()

    pdf.set_text_color(*_NAVY)
    pdf.set_font("Helvetica", "B", 14)
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(0, 7, "LAPORAN PENANGANAN ISU")
    pdf.set_font("Helvetica", size=9)
    pdf.set_text_color(*_MUTED)
    pdf.set_x(pdf.l_margin)
    pdf.cell(0, 5, _safe(f"Nomor: {doc_no}"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_x(pdf.l_margin)
    pdf.cell(0, 5, _safe(f"Dicetak: {printed}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    _brief_heading(pdf, "I", "Identitas isu")
    _brief_label_value(pdf, "Judul", issue.title or "-")
    _brief_label_value(pdf, "Tingkat risiko", risk_label)
    _brief_label_value(pdf, "Status", f"{status_label}  |  Penanggung jawab: {actor}")
    _brief_label_value(pdf, "Sumber", outlet)
    if source_url:
        _brief_label_value(pdf, "Tautan", source_url)
    _brief_label_value(pdf, "Tercatat", _fmt_wib(issue.created_at))
    if issue.status == "closed" and issue.closed_at:
        _brief_label_value(pdf, "Ditutup", _fmt_wib(issue.closed_at))

    _brief_heading(pdf, "II", "Ringkasan")
    _brief_paragraph(pdf, summary or "Ringkasan belum tersedia.")
    if why:
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_text_color(*_MUTED)
        pdf.cell(0, 5, "Konteks terkini", new_x="LMARGIN", new_y="NEXT")
        _brief_paragraph(pdf, why)
    if statement:
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_text_color(*_MUTED)
        pdf.cell(0, 5, "Narasi resmi", new_x="LMARGIN", new_y="NEXT")
        _brief_paragraph(pdf, statement)

    _brief_heading(pdf, "III", "Sumber dan bukti")
    if not evidence:
        _brief_paragraph(pdf, "Belum ada bukti terlampir.")
    else:
        for idx, item in enumerate(evidence[:8], start=1):
            name = (item.source_name or item.title or "Sumber").strip()
            pdf.set_x(pdf.l_margin)
            pdf.set_font("Helvetica", "B", 10)
            pdf.set_text_color(*_INK)
            pdf.multi_cell(0, 5.2, _safe(f"{idx}. {name}"))
            snippet = (item.snippet or "").strip()
            if snippet:
                pdf.set_x(pdf.l_margin + 4)
                pdf.set_font("Helvetica", size=9)
                pdf.set_text_color(*_INK)
                pdf.multi_cell(0, 4.8, _safe(snippet[:400]))
            if item.url:
                pdf.set_x(pdf.l_margin + 4)
                pdf.set_font("Helvetica", size=8)
                pdf.set_text_color(*_MUTED)
                pdf.multi_cell(0, 4.4, _safe(item.url))
            pdf.ln(1)

    _brief_heading(pdf, "IV", "Tindak lanjut")
    _brief_label_value(
        pdf,
        "Validasi OPD",
        f"{sum(1 for v in vals if v.status == 'validated')} dari {len(vals)} selesai",
    )
    if vals:
        for idx, item in enumerate(vals[:8], start=1):
            label = _VALIDATION_LABEL.get(item.status, item.status or "-")
            note = (item.response_notes or item.request_notes or "").strip()
            line = f"{idx}. {item.opd_name} - {label}"
            if note:
                line = f"{line}. {note[:180]}"
            pdf.set_x(pdf.l_margin)
            pdf.set_font("Helvetica", size=10)
            pdf.set_text_color(*_INK)
            pdf.multi_cell(0, 5.2, _safe(line))
    else:
        _brief_paragraph(pdf, "Belum ada permintaan validasi kepada OPD.")

    pdf.ln(1)
    _brief_label_value(pdf, "Konten", f"{len(contents)} produk")
    if contents:
        for idx, item in enumerate(contents[:8], start=1):
            pdf.set_x(pdf.l_margin)
            pdf.set_font("Helvetica", size=10)
            pdf.set_text_color(*_INK)
            kind = _lbl(_CONTENT_TYPE, item.content_type, "Konten")
            state = _lbl(_CONTENT_STATUS, item.status)
            pdf.multi_cell(0, 5.2, _safe(f"{idx}. {item.title} — {kind}, {state}"))
    _brief_label_value(pdf, "Media blast", f"{blasts} pengiriman")

    if actions:
        _brief_heading(pdf, "V", "Rekomendasi tindakan")
        for idx, action in enumerate([a for a in actions if a][:6], start=1):
            pdf.set_x(pdf.l_margin)
            pdf.set_font("Helvetica", size=10)
            pdf.set_text_color(*_INK)
            pdf.multi_cell(0, 5.2, _safe(f"{idx}. {action}"))

    pdf.ln(4)
    pdf.set_x(pdf.l_margin)
    pdf.set_font("Helvetica", "I", 9)
    pdf.set_text_color(*_MUTED)
    pdf.multi_cell(
        0,
        4.8,
        _safe(
            "Demikian laporan ini disusun sebagai bahan koordinasi penanganan isu "
            "di lingkungan Pemerintah Provinsi Banten."
        ),
    )

    out = pdf.output()
    pdf_bytes = bytes(out) if isinstance(out, (bytes, bytearray)) else out.encode("latin-1")
    return pdf_bytes, "application/pdf", f"laporan-isu-{issue.id}.pdf"


def _crisis_cover(data: dict, user_id: int | None) -> dict:
    issues = data["issues"]
    critical = sum(1 for i in issues if i["risk_level"] in {"R3", "R4", "R5"})
    open_count = sum(1 for i in issues if i["status"] != "closed")
    year = _now().year
    narrative = (
        f"Laporan ini merangkum { _fmt_int(data['count']) } isu yang tercatat pada SIAGAPIM. "
        f"Sebanyak { _fmt_int(open_count) } isu masih dalam penanganan, "
        f"dan { _fmt_int(critical) } di antaranya berada pada tingkat risiko meningkat hingga kritis. "
        "Rincian sumber, progres validasi organisasi perangkat daerah, produk konten, "
        "dan penyebarluasan media dimuat pada lembar berikutnya."
    )
    return {
        "title": "LAPORAN PENANGANAN KRISIS",
        "doc_no": f"SIAGAPIM/KRISIS/{year}",
        "printed_at": _fmt_wib(_now()),
        "printed_by": _printer_name(user_id),
        "subject": "Ikhtisar penanganan isu dan progres koordinasi",
        "facts": [
            ("Jumlah isu", _fmt_int(data["count"])),
            ("Masih ditangani", _fmt_int(open_count)),
            ("Risiko meningkat ke atas", _fmt_int(critical)),
            ("Validasi selesai", _fmt_int(sum(i["validated"] for i in issues))),
            ("Produk konten", _fmt_int(sum(i["contents"] for i in issues))),
            ("Pengiriman media", _fmt_int(sum(i["blasts"] for i in issues))),
        ],
        "narrative": narrative,
    }


def export_crisis(fmt: str, issue_id: int | None, user_id: int | None, project_id: str | None = None):
    if fmt == "pdf" and issue_id:
        return _export_issue_brief(issue_id, user_id, project_id)
    data = collect_crisis(issue_id, project_id=project_id)
    title = "Laporan Penanganan Krisis"
    _save_report("crisis", title, data, user_id)
    cover = _crisis_cover(data, user_id)
    issue_rows = [
        [
            idx,
            i["title"],
            i["risk_label"],
            i["status_label"],
            i["actor"],
            i["source"],
            i["created_label"],
            f"{i['validated']} dari {i['validations']}",
            f"{i['approved']} dari {i['contents']}",
            i["blasts"],
            (i["summary"] or "-")[:280],
        ]
        for idx, i in enumerate(data["issues"], start=1)
    ]
    if fmt == "xlsx":
        blob = _formal_workbook(
            cover,
            [
                {
                    "name": "Daftar Isu",
                    "headers": [
                        "No.",
                        "Judul isu",
                        "Tingkat risiko",
                        "Status",
                        "Penanggung jawab",
                        "Sumber",
                        "Tercatat",
                        "Validasi OPD",
                        "Konten disetujui",
                        "Blast",
                        "Ringkasan",
                    ],
                    "rows": issue_rows,
                    "widths": [6, 42, 22, 22, 18, 22, 28, 16, 18, 10, 46],
                }
            ],
        )
        return blob, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "laporan-krisis.xlsx"
    shown = data["issues"][:20]
    blocks = [
        ("h", "I", "Ikhtisar"),
        ("p", cover["narrative"]),
        ("kv", cover["facts"]),
        ("h", "II", "Daftar isu"),
        (
            "table",
            ["No.", "Risiko", "Status", "Judul"],
            [[n, i["risk_label"], i["status_label"], i["title"]] for n, i in enumerate(shown, start=1)],
            (8, 28, 32, 50),
        ),
        ("h", "III", "Uraian penanganan"),
        (
            "items",
            [
                (
                    f"{n}. {i['title']}. Risiko {i['risk_label']}, status {i['status_label'].lower()} "
                    f"(penanggung jawab: {i['actor']}). Sumber: {i['source']}. "
                    f"Tercatat {i['created_label']}. Validasi {i['validated']} dari {i['validations']}, "
                    f"konten disetujui {i['approved']} dari {i['contents']}, blast {i['blasts']}. "
                    f"{(i['summary'] or 'Ringkasan belum tersedia.')[:320]}"
                )
                for n, i in enumerate(shown, start=1)
            ],
        ),
        (
            "close",
            "Demikian laporan ini disusun sebagai bahan koordinasi penanganan isu "
            "di lingkungan Pemerintah Provinsi Banten. Rincian lengkap terlampir pada berkas Excel.",
        ),
    ]
    pdf_bytes = _official_pdf("Laporan Penanganan Krisis", "LAPORAN PENANGANAN KRISIS", cover["doc_no"], cover["printed_by"], blocks)
    return pdf_bytes, "application/pdf", "laporan-krisis.pdf"


# ── Media SLA ───────────────────────────────────────────────────────────────


def collect_media_sla() -> dict:
    logs = MediaSlaLog.query.order_by(MediaSlaLog.created_at.desc()).limit(200).all()
    partners = {p.id: p for p in MediaPartner.query.all()}
    rows = []
    for log in logs:
        p = partners.get(log.media_partner_id)
        rows.append(
            {
                "id": log.id,
                "media": p.name if p else f"Mitra #{log.media_partner_id}",
                "editor": (p.editor_name if p else None) or "-",
                "area": (p.coverage_area if p else None) or "-",
                "response_minutes": log.response_minutes,
                "sla_compliant": bool(log.sla_compliant),
                "content_match": bool(log.content_match),
                "published_at": log.published_at.isoformat() if log.published_at else None,
                "published_label": _fmt_wib(log.published_at),
                "notes": (log.notes or "").strip(),
            }
        )
    total = len(rows)
    compliant = sum(1 for r in rows if r["sla_compliant"])
    minutes = [r["response_minutes"] for r in rows if r["response_minutes"] is not None]
    by_media: dict[str, dict] = {}
    for row in rows:
        bucket = by_media.setdefault(
            row["media"],
            {"media": row["media"], "editor": row["editor"], "area": row["area"], "total": 0, "ok": 0, "minutes": []},
        )
        bucket["total"] += 1
        bucket["ok"] += 1 if row["sla_compliant"] else 0
        if row["response_minutes"] is not None:
            bucket["minutes"].append(row["response_minutes"])
    ranking = []
    for bucket in by_media.values():
        avg = round(sum(bucket["minutes"]) / len(bucket["minutes"])) if bucket["minutes"] else None
        ranking.append({**bucket, "avg": avg, "rate": _fmt_pct(bucket["ok"], bucket["total"])})
    ranking.sort(key=lambda item: (item["ok"] / item["total"] if item["total"] else 0, -(item["avg"] or 999)))
    return {
        "type": "media_sla",
        "total": total,
        "compliant": compliant,
        "rate": round(100 * compliant / total, 1) if total else 0,
        "rate_label": _fmt_pct(compliant, total),
        "avg_minutes": round(sum(minutes) / len(minutes)) if minutes else None,
        "logs": rows,
        "ranking": ranking,
    }


def export_media_sla(fmt: str, user_id: int | None):
    data = collect_media_sla()
    title = "Laporan Aktivitas Media (SLA)"
    _save_report("media_sla", title, data, user_id)
    year = _now().year
    late = data["total"] - data["compliant"]
    avg = f"{data['avg_minutes']} menit" if data["avg_minutes"] is not None else "belum dapat dihitung"
    narrative = (
        f"Dari { _fmt_int(data['total']) } catatan tayang mitra media, "
        f"{ _fmt_int(data['compliant']) } memenuhi batas waktu 60 menit ({data['rate_label']}). "
        f"Rata-rata waktu respons {avg}. Sebanyak { _fmt_int(late) } tayang melewati batas "
        "dan perlu ditindaklanjuti kepada pemred mitra yang bersangkutan."
    )
    cover = {
        "title": "LAPORAN AKTIVITAS MEDIA",
        "doc_no": f"SIAGAPIM/SLA/{year}",
        "printed_at": _fmt_wib(_now()),
        "printed_by": _printer_name(user_id),
        "subject": "Kepatuhan waktu tayang mitra media terhadap batas 60 menit",
        "facts": [
            ("Catatan tayang", _fmt_int(data["total"])),
            ("Memenuhi SLA", f"{_fmt_int(data['compliant'])} ({data['rate_label']})"),
            ("Melewati SLA", _fmt_int(late)),
            ("Rata-rata respons", avg),
            ("Batas waktu", "60 menit sejak pengiriman"),
            ("Jumlah mitra tercatat", _fmt_int(len(data["ranking"]))),
        ],
        "narrative": narrative,
    }
    rank_rows = [
        [n, r["media"], r["editor"], r["area"], r["total"], r["ok"], r["total"] - r["ok"], r["rate"], r["avg"] if r["avg"] is not None else "-"]
        for n, r in enumerate(data["ranking"], start=1)
    ]
    log_rows = [
        [
            n,
            r["media"],
            r["editor"],
            r["published_label"],
            r["response_minutes"] if r["response_minutes"] is not None else "-",
            "Memenuhi" if r["sla_compliant"] else "Melewati batas",
            "Sesuai" if r["content_match"] else "Tidak sesuai",
            r["notes"] or "-",
        ]
        for n, r in enumerate(data["logs"], start=1)
    ]
    if fmt == "xlsx":
        blob = _formal_workbook(
            cover,
            [
                {
                    "name": "Peringkat Mitra",
                    "headers": ["No.", "Mitra media", "Pemimpin redaksi", "Wilayah", "Tayang", "Memenuhi", "Lewat", "Kepatuhan", "Rata-rata menit"],
                    "rows": rank_rows,
                    "widths": [6, 28, 24, 22, 12, 14, 12, 14, 18],
                },
                {
                    "name": "Riwayat Tayang",
                    "headers": ["No.", "Mitra media", "Pemimpin redaksi", "Waktu tayang", "Menit", "SLA", "Kesesuaian naskah", "Catatan"],
                    "rows": log_rows,
                    "widths": [6, 26, 22, 28, 12, 18, 20, 40],
                },
            ],
        )
        return blob, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "laporan-sla-media.xlsx"
    late_items = [
        f"{r['media']} menayangkan dalam {r['response_minutes'] or '-'} menit pada {r['published_label']}. {r['notes']}".strip()
        for r in data["logs"]
        if not r["sla_compliant"]
    ][:12]
    blocks = [
        ("h", "I", "Ikhtisar kepatuhan"),
        ("p", narrative),
        ("kv", cover["facts"]),
        ("h", "II", "Peringkat mitra"),
        (
            "table",
            ["No.", "Mitra", "Tayang", "Kepatuhan", "Rata-rata"],
            [[n, r["media"], r["total"], r["rate"], f"{r['avg']} mnt" if r["avg"] is not None else "-"] for n, r in enumerate(data["ranking"], start=1)],
            (8, 42, 16, 20, 22),
        ),
        ("h", "III", "Tayang yang melewati batas"),
        ("items", [f"{n}. {line}" for n, line in enumerate(late_items, start=1)]),
        (
            "close",
            "Demikian laporan aktivitas media disusun untuk bahan evaluasi kemitraan "
            "dan penegakan batas waktu tayang di lingkungan Pemerintah Provinsi Banten.",
        ),
    ]
    return (
        _official_pdf("Laporan Aktivitas Media", "LAPORAN AKTIVITAS MEDIA", cover["doc_no"], cover["printed_by"], blocks),
        "application/pdf",
        "laporan-sla-media.pdf",
    )


# ── ASN ─────────────────────────────────────────────────────────────────────


def collect_asn() -> dict:
    rows = (
        db.session.query(
            MissionParticipation.opd_name,
            func.count(MissionParticipation.id),
        )
        .group_by(MissionParticipation.opd_name)
        .order_by(func.count(MissionParticipation.id).desc())
        .all()
    )
    grand = sum(r[1] for r in rows)
    by_opd = [
        {"opd_name": r[0] or "Belum mencantumkan OPD", "total": r[1], "share": _fmt_pct(r[1], grand)}
        for r in rows
    ]
    missions = Mission.query.order_by(Mission.created_at.desc()).limit(50).all()
    mission_rows = []
    detail = []
    for mission in missions:
        done = len(mission.participations)
        target = mission.target_count or 0
        mission_rows.append(
            {
                "id": mission.id,
                "title": mission.title,
                "status": mission.status,
                "status_label": _lbl(_MISSION_STATUS, mission.status),
                "action": _lbl(_ACTION_LABEL, mission.action_type, "Aksi organik"),
                "participations": done,
                "target_count": target,
                "progress": _fmt_pct(done, target) if target else "-",
                "instruction": (mission.instruction or "").strip(),
                "period": f"{_fmt_wib(mission.starts_at)} s.d. {_fmt_wib(mission.ends_at)}"
                if mission.starts_at or mission.ends_at
                else "Tanpa batas waktu",
            }
        )
        for item in mission.participations:
            detail.append(
                {
                    "mission": mission.title,
                    "opd": item.opd_name or "Belum mencantumkan OPD",
                    "notes": (item.notes or "").strip() or "-",
                    "at": _fmt_wib(item.completed_at),
                }
            )
    active = sum(1 for m in mission_rows if m["status"] == "active")
    return {
        "type": "asn",
        "grand_total": grand,
        "active_missions": active,
        "by_opd": by_opd,
        "missions": mission_rows,
        "detail": detail,
    }


def export_asn(fmt: str, user_id: int | None):
    data = collect_asn()
    title = "Rekap Partisipasi ASN per OPD"
    _save_report("asn", title, data, user_id)
    year = _now().year
    narrative = (
        f"Tercatat { _fmt_int(data['grand_total']) } partisipasi aparatur sipil negara "
        f"dari { _fmt_int(len(data['by_opd'])) } organisasi perangkat daerah. "
        f"Sebanyak { _fmt_int(data['active_missions']) } misi masih berjalan dari "
        f"{ _fmt_int(len(data['missions'])) } misi yang dilaporkan. "
        "Angka ini dipakai untuk menilai pemerataan amplifikasi organik antarperangkat daerah."
    )
    cover = {
        "title": "REKAP PARTISIPASI APARATUR SIPIL NEGARA",
        "doc_no": f"SIAGAPIM/ASN/{year}",
        "printed_at": _fmt_wib(_now()),
        "printed_by": _printer_name(user_id),
        "subject": "Partisipasi ASN per organisasi perangkat daerah pada mission board",
        "facts": [
            ("Total partisipasi", _fmt_int(data["grand_total"])),
            ("Organisasi perangkat daerah", _fmt_int(len(data["by_opd"]))),
            ("Misi dilaporkan", _fmt_int(len(data["missions"]))),
            ("Misi masih berjalan", _fmt_int(data["active_missions"])),
        ],
        "narrative": narrative,
    }
    if fmt == "xlsx":
        blob = _formal_workbook(
            cover,
            [
                {
                    "name": "Per OPD",
                    "headers": ["No.", "Organisasi perangkat daerah", "Partisipasi", "Porsi"],
                    "rows": [[n, r["opd_name"], r["total"], r["share"]] for n, r in enumerate(data["by_opd"], start=1)],
                    "widths": [6, 42, 16, 14],
                },
                {
                    "name": "Misi",
                    "headers": ["No.", "Judul misi", "Status", "Jenis aksi", "Partisipasi", "Target", "Capaian", "Periode", "Instruksi"],
                    "rows": [
                        [n, m["title"], m["status_label"], m["action"], m["participations"], m["target_count"] or "-", m["progress"], m["period"], m["instruction"][:240] or "-"]
                        for n, m in enumerate(data["missions"], start=1)
                    ],
                    "widths": [6, 36, 16, 28, 14, 12, 12, 36, 46],
                },
                {
                    "name": "Rincian Partisipasi",
                    "headers": ["No.", "Misi", "Organisasi perangkat daerah", "Waktu", "Catatan"],
                    "rows": [[n, d["mission"], d["opd"], d["at"], d["notes"]] for n, d in enumerate(data["detail"], start=1)],
                    "widths": [6, 40, 32, 28, 40],
                },
            ],
        )
        return blob, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "rekap-asn.xlsx"
    blocks = [
        ("h", "I", "Ikhtisar"),
        ("p", narrative),
        ("kv", cover["facts"]),
        ("h", "II", "Partisipasi per organisasi perangkat daerah"),
        (
            "table",
            ["No.", "Organisasi perangkat daerah", "Jumlah", "Porsi"],
            [[n, r["opd_name"], r["total"], r["share"]] for n, r in enumerate(data["by_opd"], start=1)],
            (8, 55, 18, 16),
        ),
        ("h", "III", "Progres misi"),
        (
            "items",
            [
                (
                    f"{n}. {m['title']} — {m['status_label'].lower()}. "
                    f"Aksi: {m['action'].lower()}. Capaian {m['participations']} dari {m['target_count'] or 'tanpa'} target "
                    f"({m['progress']}). {m['instruction'][:220]}"
                )
                for n, m in enumerate(data["missions"][:15], start=1)
            ],
        ),
        (
            "close",
            "Demikian rekap partisipasi disusun sebagai bahan evaluasi keterlibatan "
            "aparatur sipil negara dalam amplifikasi informasi resmi Pemerintah Provinsi Banten.",
        ),
    ]
    return (
        _official_pdf("Rekap Partisipasi ASN", "REKAP PARTISIPASI APARATUR SIPIL NEGARA", cover["doc_no"], cover["printed_by"], blocks),
        "application/pdf",
        "rekap-asn.pdf",
    )


# ── KOL ─────────────────────────────────────────────────────────────────────


def collect_kol() -> dict:
    campaigns = KolCampaign.query.order_by(KolCampaign.created_at.desc()).limit(100).all()
    partners = KolPartner.query.order_by(KolPartner.name.asc()).all()
    by_id = {p.id: p for p in partners}
    counts: dict[int, int] = {}
    rows = []
    for campaign in campaigns:
        partner = by_id.get(campaign.kol_id)
        counts[campaign.kol_id] = counts.get(campaign.kol_id, 0) + 1
        rows.append(
            {
                "id": campaign.id,
                "title": campaign.title,
                "kol": partner.name if partner else f"Mitra #{campaign.kol_id}",
                "platform": (partner.platform if partner else None) or "-",
                "handle": (partner.handle if partner else None) or "-",
                "status": campaign.status,
                "status_label": _lbl(_CAMPAIGN_STATUS, campaign.status),
                "views": campaign.views or 0,
                "likes": campaign.likes or 0,
                "comments": campaign.comments or 0,
                "budget": float(campaign.budget) if campaign.budget is not None else None,
                "budget_label": _fmt_idr(campaign.budget),
                "budget_status": _lbl(_BUDGET_LABEL, campaign.budget_status),
                "published_label": _fmt_wib(campaign.published_at),
                "deliverable_url": campaign.deliverable_url or "-",
                "notes": (campaign.notes or "").strip(),
            }
        )
    partner_rows = [
        {
            "name": p.name,
            "platform": p.platform or "-",
            "handle": p.handle or "-",
            "followers": p.followers or 0,
            "engagement": f"{(p.engagement_rate or 0):.1f}".replace(".", ",") + "%",
            "contract": _lbl(_CONTRACT_LABEL, p.contract_status),
            "topics": p.topics or "-",
            "campaigns": counts.get(p.id, 0),
            "active": "Aktif" if p.is_active else "Nonaktif",
        }
        for p in partners
    ]
    return {
        "type": "kol",
        "count": len(rows),
        "total_views": sum(r["views"] for r in rows),
        "total_likes": sum(r["likes"] for r in rows),
        "total_comments": sum(r["comments"] for r in rows),
        "total_budget": sum(r["budget"] or 0 for r in rows),
        "campaigns": rows,
        "partners": partner_rows,
    }


def export_kol(fmt: str, user_id: int | None):
    data = collect_kol()
    title = "Laporan Kinerja KOL dan Kampanye"
    _save_report("kol", title, data, user_id)
    year = _now().year
    narrative = (
        f"Terdapat { _fmt_int(len(data['partners'])) } mitra key opinion leader dan "
        f"{ _fmt_int(data['count']) } kampanye. Total jangkauan tercatat "
        f"{ _fmt_int(data['total_views']) } tayangan, { _fmt_int(data['total_likes']) } suka, "
        f"dan { _fmt_int(data['total_comments']) } komentar. Nilai anggaran yang tercantum "
        f"sebesar { _fmt_idr(data['total_budget']) }."
    )
    cover = {
        "title": "LAPORAN KINERJA KEY OPINION LEADER",
        "doc_no": f"SIAGAPIM/KOL/{year}",
        "printed_at": _fmt_wib(_now()),
        "printed_by": _printer_name(user_id),
        "subject": "Kinerja mitra dan kampanye amplifikasi",
        "facts": [
            ("Mitra", _fmt_int(len(data["partners"]))),
            ("Kampanye", _fmt_int(data["count"])),
            ("Tayangan", _fmt_int(data["total_views"])),
            ("Suka", _fmt_int(data["total_likes"])),
            ("Komentar", _fmt_int(data["total_comments"])),
            ("Nilai anggaran", _fmt_idr(data["total_budget"])),
        ],
        "narrative": narrative,
    }
    if fmt == "xlsx":
        blob = _formal_workbook(
            cover,
            [
                {
                    "name": "Mitra",
                    "headers": ["No.", "Nama", "Platform", "Akun", "Pengikut", "Engagement", "Kontrak", "Topik", "Kampanye", "Status mitra"],
                    "rows": [
                        [n, p["name"], p["platform"], p["handle"], p["followers"], p["engagement"], p["contract"], p["topics"], p["campaigns"], p["active"]]
                        for n, p in enumerate(data["partners"], start=1)
                    ],
                    "widths": [6, 24, 14, 18, 14, 14, 20, 28, 12, 14],
                },
                {
                    "name": "Kampanye",
                    "headers": ["No.", "Judul", "Mitra", "Platform", "Status", "Tayangan", "Suka", "Komentar", "Anggaran", "Status anggaran", "Tayang pada", "Tautan"],
                    "rows": [
                        [n, c["title"], c["kol"], c["platform"], c["status_label"], c["views"], c["likes"], c["comments"], c["budget_label"], c["budget_status"], c["published_label"], c["deliverable_url"]]
                        for n, c in enumerate(data["campaigns"], start=1)
                    ],
                    "widths": [6, 36, 22, 14, 16, 12, 12, 12, 18, 20, 28, 36],
                },
            ],
        )
        return blob, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "laporan-kol.xlsx"
    blocks = [
        ("h", "I", "Ikhtisar"),
        ("p", narrative),
        ("kv", cover["facts"]),
        ("h", "II", "Mitra"),
        (
            "table",
            ["No.", "Nama", "Platform", "Kontrak", "Kampanye"],
            [[n, p["name"], p["platform"], p["contract"], p["campaigns"]] for n, p in enumerate(data["partners"], start=1)],
            (8, 40, 22, 28, 16),
        ),
        ("h", "III", "Kampanye"),
        (
            "items",
            [
                (
                    f"{n}. {c['title']} oleh {c['kol']} ({c['platform']}). Status {c['status_label'].lower()}. "
                    f"Tayangan {_fmt_int(c['views'])}, suka {_fmt_int(c['likes'])}, komentar {_fmt_int(c['comments'])}. "
                    f"Anggaran {c['budget_label']} ({c['budget_status'].lower()})."
                )
                for n, c in enumerate(data["campaigns"][:20], start=1)
            ],
        ),
        (
            "close",
            "Demikian laporan kinerja key opinion leader disusun sebagai bahan evaluasi "
            "amplifikasi dan pertanggungjawaban anggaran di lingkungan Pemerintah Provinsi Banten.",
        ),
    ]
    return (
        _official_pdf("Laporan Kinerja KOL", "LAPORAN KINERJA KEY OPINION LEADER", cover["doc_no"], cover["printed_by"], blocks),
        "application/pdf",
        "laporan-kol.pdf",
    )


# ── Executive ───────────────────────────────────────────────────────────────


def collect_executive(project_id: str | None = None) -> dict:
    from app.utils.project_scope import filter_issues_query, filter_query_by_issue_ids, issue_ids_for_project

    open_statuses = {"open", "validating", "producing", "approved"}
    issues = filter_issues_query(Issue.query, project_id).all()
    active = [i for i in issues if i.status in open_statuses]
    critical = [i for i in active if i.risk_level in {"R3", "R4", "R5"}]
    waiting_q = OpdValidation.query.filter_by(status="waiting")
    waiting_q = filter_query_by_issue_ids(waiting_q, OpdValidation.issue_id, project_id)
    waiting = waiting_q.count()
    waiting_rows = waiting_q.order_by(OpdValidation.requested_at.asc()).limit(30).all()
    issue_titles = {i.id: i.title for i in issues}
    waiting_items = [
        {
            "opd": v.opd_name,
            "issue": issue_titles.get(v.issue_id, f"Isu #{v.issue_id}"),
            "requested": _fmt_wib(v.requested_at),
            "notes": (v.request_notes or "").strip() or "-",
        }
        for v in waiting_rows
    ]
    ids = issue_ids_for_project(project_id)
    asn_q = MissionParticipation.query.join(Mission, Mission.id == MissionParticipation.mission_id)
    if ids is not None:
        asn_q = asn_q.filter(Mission.issue_id.in_(ids or [-1]))
    asn_total = asn_q.count()
    kol_q = db.session.query(func.coalesce(func.sum(KolCampaign.views), 0))
    if ids is not None:
        kol_q = kol_q.filter(KolCampaign.issue_id.in_(ids or [-1]))
    kol_views = kol_q.scalar() or 0
    return {
        "type": "executive",
        "headline": f"{len(critical)} isu kritis aktif dari {len(active)} isu terbuka",
        "kpis": {
            "active_issues": len(active),
            "critical_issues": len(critical),
            "waiting_validations": waiting,
            "asn_participations": asn_total,
            "kol_views": int(kol_views),
        },
        "critical_issues": [
            {
                "title": i.title,
                "risk": _lbl(_RISK_LABEL, i.risk_level),
                "status": _STATUS_LABEL.get(i.status, (i.status or "-", "-"))[0],
                "source": i.primary_source_label() or "Tidak tercatat",
                "summary": (i.summary or "").strip()[:280] or "Ringkasan belum tersedia.",
                "recorded": _fmt_wib(i.created_at),
            }
            for i in critical[:12]
        ],
        "waiting_items": waiting_items,
        "recommendations": [],
    }


def export_executive(fmt: str, user_id: int | None, project_id: str | None = None):
    data = collect_executive(project_id=project_id)
    title = "Executive Brief"
    k = data["kpis"]
    recommendations = []
    if k["critical_issues"]:
        recommendations.append(
            f"Segera koordinasikan { _fmt_int(k['critical_issues']) } isu berisiko meningkat hingga kritis kepada penanggung jawab masing-masing."
        )
    if k["waiting_validations"]:
        recommendations.append(
            f"Tindaklanjuti { _fmt_int(k['waiting_validations']) } permintaan validasi yang belum dijawab organisasi perangkat daerah."
        )
    if k["active_issues"] and not k["critical_issues"]:
        recommendations.append("Pertahankan pemantauan isu yang masih terbuka agar tidak naik tingkat risikonya.")
    recommendations.append("Pastikan naskah yang telah disetujui segera disebarluaskan melalui mitra media dan kanal resmi.")
    if k["asn_participations"] == 0:
        recommendations.append("Aktifkan misi aparatur sipil negara untuk amplifikasi organik pada isu prioritas.")
    data["recommendations"] = recommendations
    _save_report("executive", title, data, user_id)
    year = _now().year
    narrative = (
        f"Pada saat laporan disusun, terdapat { _fmt_int(k['active_issues']) } isu yang masih ditangani, "
        f"termasuk { _fmt_int(k['critical_issues']) } isu berisiko meningkat hingga kritis. "
        f"Validasi yang belum dijawab berjumlah { _fmt_int(k['waiting_validations']) }. "
        f"Partisipasi aparatur tercatat { _fmt_int(k['asn_participations']) } kali, "
        f"dengan jangkauan key opinion leader sebesar { _fmt_int(k['kol_views']) } tayangan."
    )
    cover = {
        "title": "RINGKASAN EKSEKUTIF PENANGANAN ISU",
        "doc_no": f"SIAGAPIM/EKSEKUTIF/{year}",
        "printed_at": _fmt_wib(_now()),
        "printed_by": _printer_name(user_id),
        "subject": "Bahan pimpinan mengenai isu aktif, validasi, dan amplifikasi",
        "facts": [
            ("Isu masih ditangani", _fmt_int(k["active_issues"])),
            ("Isu berisiko tinggi", _fmt_int(k["critical_issues"])),
            ("Validasi belum dijawab", _fmt_int(k["waiting_validations"])),
            ("Partisipasi ASN", _fmt_int(k["asn_participations"])),
            ("Tayangan KOL", _fmt_int(k["kol_views"])),
        ],
        "narrative": narrative,
    }
    critical_rows = [
        [n, i["title"], i["risk"], i["status"], i["source"], i["recorded"], i["summary"]]
        for n, i in enumerate(data["critical_issues"], start=1)
    ]
    waiting_rows = [
        [n, item["opd"], item["issue"], item["requested"], item["notes"]]
        for n, item in enumerate(data["waiting_items"], start=1)
    ]
    if fmt == "xlsx":
        blob = _formal_workbook(
            cover,
            [
                {
                    "name": "Isu Prioritas",
                    "headers": ["No.", "Judul", "Tingkat risiko", "Status", "Sumber", "Tercatat", "Ringkasan"],
                    "rows": critical_rows,
                    "widths": [6, 40, 22, 20, 22, 28, 46],
                },
                {
                    "name": "Validasi Menunggu",
                    "headers": ["No.", "Organisasi perangkat daerah", "Isu", "Diminta pada", "Catatan permintaan"],
                    "rows": waiting_rows,
                    "widths": [6, 32, 40, 28, 42],
                },
                {
                    "name": "Rekomendasi",
                    "headers": ["No.", "Rekomendasi"],
                    "rows": [[n, text] for n, text in enumerate(recommendations, start=1)],
                    "widths": [6, 90],
                },
            ],
        )
        return blob, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "executive-brief.xlsx"
    blocks = [
        ("h", "I", "Ikhtisar untuk pimpinan"),
        ("p", narrative),
        ("kv", cover["facts"]),
        ("h", "II", "Isu berisiko tinggi"),
        (
            "items",
            [
                f"{n}. {i['title']}. {i['risk']}, status {i['status'].lower()}, sumber {i['source']}, tercatat {i['recorded']}. {i['summary']}"
                for n, i in enumerate(data["critical_issues"], start=1)
            ],
        ),
        ("h", "III", "Validasi yang belum dijawab"),
        (
            "items",
            [
                f"{n}. {item['opd']} atas isu \"{item['issue']}\", diminta {item['requested']}."
                for n, item in enumerate(data["waiting_items"][:12], start=1)
            ],
        ),
        ("h", "IV", "Rekomendasi"),
        ("items", [f"{n}. {text}" for n, text in enumerate(recommendations, start=1)]),
        (
            "close",
            "Demikian ringkasan eksekutif ini disampaikan untuk bahan arahan pimpinan "
            "Pemerintah Provinsi Banten.",
        ),
    ]
    return (
        _official_pdf("Ringkasan Eksekutif", "RINGKASAN EKSEKUTIF PENANGANAN ISU", cover["doc_no"], cover["printed_by"], blocks),
        "application/pdf",
        "executive-brief.pdf",
    )
